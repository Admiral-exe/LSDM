"""
Atomic Two-Way Excel Writer.
Writes updates and deletions safely using tempfile verification, explicit string typing (@),
and atomic os.replace, with automatic lockfile checking and backup snapshots.
"""
import os
import uuid
from pathlib import Path
from typing import Dict, Any, Optional
import datetime
import openpyxl
from openpyxl.styles import numbers

import config
from backend.excel_lock import is_file_locked
from backend.excel_backup import create_backup
from backend.database import log_audit

def parse_date_parts(date_val: Any) -> tuple[Optional[int], Optional[int], Optional[int]]:
    """Parses a date string (YYYY-MM-DD or DD/MM/YYYY) or datetime object into (DD, MM, YYYY)."""
    if not date_val:
        return None, None, None
    if isinstance(date_val, (datetime.date, datetime.datetime)):
        return date_val.day, date_val.month, date_val.year
    date_str = str(date_val).strip()
    if not date_str:
        return None, None, None
    try:
        if "-" in date_str:
            parts = date_str.split("-")
            if len(parts[0]) == 4: # YYYY-MM-DD
                return int(parts[2]), int(parts[1]), int(parts[0])
            else: # DD-MM-YYYY
                return int(parts[0]), int(parts[1]), int(parts[2])
        elif "/" in date_str:
            parts = date_str.split("/")
            if len(parts[2]) == 4: # DD/MM/YYYY
                return int(parts[0]), int(parts[1]), int(parts[2])
            elif len(parts[0]) == 4: # YYYY/MM/DD
                return int(parts[2]), int(parts[1]), int(parts[0])
    except Exception:
        pass
    return None, None, None

def set_cell_value(cell, value, is_strict_string: bool = False, is_integer: bool = False):
    """Sets a cell's value and formats it appropriately."""
    if value is None:
        cell.value = None
        return

    if is_strict_string:
        cell.value = str(value).strip()
        cell.data_type = 's'
        cell.number_format = '@'
    elif is_integer:
        try:
            cell.value = int(value)
            cell.data_type = 'n'
            cell.number_format = '0'
        except (ValueError, TypeError):
            cell.value = str(value)
            cell.data_type = 's'
    else:
        cell.value = value

class AtomicExcelWriter:
    @staticmethod
    def _prepare_write(target_path: Path, action_tag: str):
        """Verifies lock and creates backup before opening."""
        locked, reason = is_file_locked(target_path)
        if locked:
            raise PermissionError(f"Target file is locked: {reason}")
        create_backup(target_path, action_tag)

    @staticmethod
    def _finalize_write(wb: openpyxl.Workbook, target_path: Path):
        """Saves to tempfile, verifies structure, and atomically swaps."""
        temp_filename = f".tmp_{uuid.uuid4().hex}.xlsx"
        temp_path = target_path.parent / temp_filename

        try:
            wb.save(temp_path)
            # Verify temporary workbook can be loaded without corruption
            verify_wb = openpyxl.load_workbook(temp_path, data_only=True)
            if not verify_wb.sheetnames:
                raise ValueError("Verification failed: Empty workbook produced.")
            verify_wb.close()

            # Atomic swap
            os.replace(temp_path, target_path)
            log_audit(
                action="ATOMIC_WRITE_SUCCESS",
                target_file=target_path.name,
                details=f"File successfully updated atomically: {target_path.name}"
            )
        except Exception as e:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass
            log_audit(
                action="ATOMIC_WRITE_FAILED",
                target_file=target_path.name,
                details=f"Write failed: {str(e)}"
            )
            raise

    @classmethod
    def write_monthly_row(cls, file_path: Path | str, sheet_name: str, row_idx: int, data: Dict[str, Any]):
        """
        Writes or updates a single row in a monthly arrival file.
        Handles the ESIS column shift (17 cols vs 18 cols).
        """
        path = Path(file_path)
        cls._prepare_write(path, f"monthly_r{row_idx}")

        wb = openpyxl.load_workbook(path)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {path.name}")
        ws = wb[sheet_name]

        is_esis = (sheet_name == "ESIS")
        # Offset: standard has unnamed group in col 1; ESIS starts SR. NO in col 1
        col_offset = 0 if not is_esis else -1

        # Col 1 (A) in standard is group code
        if not is_esis:
            set_cell_value(ws.cell(row=row_idx, column=1), data.get("group_code", sheet_name), is_strict_string=True)

        # Serial No
        set_cell_value(ws.cell(row=row_idx, column=2 + col_offset), data.get("sr_no"), is_integer=True)

        # Receive Date
        set_cell_value(ws.cell(row=row_idx, column=3 + col_offset), data.get("receive_date"))

        # Sample No (Sequential integer)
        set_cell_value(ws.cell(row=row_idx, column=4 + col_offset), data.get("sample_no"), is_integer=True)

        # Division, T.P./Proce, Catgo Name, Catgo No, G/OG
        set_cell_value(ws.cell(row=row_idx, column=5 + col_offset), data.get("division", ""))
        set_cell_value(ws.cell(row=row_idx, column=6 + col_offset), data.get("tp_proce", ""))
        set_cell_value(ws.cell(row=row_idx, column=7 + col_offset), data.get("catgo_name", ""))
        set_cell_value(ws.cell(row=row_idx, column=8 + col_offset), data.get("catego_no", ""))
        set_cell_value(ws.cell(row=row_idx, column=9 + col_offset), data.get("g_og", ""))

        # Sample Name
        set_cell_value(ws.cell(row=row_idx, column=10 + col_offset), data.get("sample_name", ""))

        # Batch No (Strict String @)
        set_cell_value(ws.cell(row=row_idx, column=11 + col_offset), data.get("batch_no", ""), is_strict_string=True)

        # Exp Date, Name of Mfg
        set_cell_value(ws.cell(row=row_idx, column=12 + col_offset), data.get("exp_date", ""))
        set_cell_value(ws.cell(row=row_idx, column=13 + col_offset), data.get("name_of_mfg", ""))

        # DI Name (City)
        di_name_city = data.get("di_name_city")
        if not di_name_city:
            inspector = data.get("inspector_name", "")
            city = data.get("city_district", "")
            di_name_city = f"{inspector} ({city})" if city else inspector
        set_cell_value(ws.cell(row=row_idx, column=14 + col_offset), di_name_city)

        # Memorandum (Strict String @)
        set_cell_value(ws.cell(row=row_idx, column=15 + col_offset), data.get("memorandum", ""), is_strict_string=True)

        # Distributed To, Distribution Date, Weightage
        set_cell_value(ws.cell(row=row_idx, column=16 + col_offset), data.get("distributed_to", ""))
        set_cell_value(ws.cell(row=row_idx, column=17 + col_offset), data.get("distribution_date", ""))
        set_cell_value(ws.cell(row=row_idx, column=18 + col_offset), data.get("weightage", ""))

        cls._finalize_write(wb, path)

    @classmethod
    def write_master_row(cls, file_path: Path | str, sheet_name: str, row_idx: int, data: Dict[str, Any]):
        """
        Writes or updates a single row in the Master file.
        Handles split dates, TESTREPO multi-columns, numeric RESULT code, and text FEES.
        """
        path = Path(file_path)
        cls._prepare_write(path, f"master_r{row_idx}")

        wb = openpyxl.load_workbook(path)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {path.name}")
        ws = wb[sheet_name]

        # Col 1: SAMPNO (Strict Text @)
        set_cell_value(ws.cell(row=row_idx, column=1), data.get("sampno", ""), is_strict_string=True)

        # Cols 2, 3, 4: RECDATE (DD, MM, YYYY)
        r_dd, r_mm, r_yy = parse_date_parts(data.get("recdate_iso"))
        set_cell_value(ws.cell(row=row_idx, column=2), r_dd, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=3), r_mm, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=4), r_yy, is_integer=True)

        # Cols 5, 6, 7: SAMPDATE (DD, MM, YYYY)
        s_dd, s_mm, s_yy = parse_date_parts(data.get("sampdate_iso"))
        set_cell_value(ws.cell(row=row_idx, column=5), s_dd, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=6), s_mm, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=7), s_yy, is_integer=True)

        # Col 8: TYPE, Col 9: NAMECD, Col 10: NAME / SENDER NAME
        set_cell_value(ws.cell(row=row_idx, column=8), data.get("type", ""))
        set_cell_value(ws.cell(row=row_idx, column=9), data.get("namecd", ""))
        set_cell_value(ws.cell(row=row_idx, column=10), data.get("sender_name", ""))

        # Col 11: DISTOCD, Col 12: DISTONAME
        set_cell_value(ws.cell(row=row_idx, column=11), data.get("distocd"), is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=12), data.get("distoname", ""))

        # Col 13: PLACECD, Col 14: PLACENM
        set_cell_value(ws.cell(row=row_idx, column=13), data.get("placecd", ""))
        set_cell_value(ws.cell(row=row_idx, column=14), data.get("placenm", ""))

        # Col 15: SAMPNAME, Col 16: BATCHNO (Strict Text @), Col 17: LICENCE (Strict Text @)
        set_cell_value(ws.cell(row=row_idx, column=15), data.get("sampname", ""))
        set_cell_value(ws.cell(row=row_idx, column=16), data.get("batchno", ""), is_strict_string=True)
        set_cell_value(ws.cell(row=row_idx, column=17), data.get("licence", ""), is_strict_string=True)

        # Col 18: Preserved Empty
        set_cell_value(ws.cell(row=row_idx, column=18), None)

        # Cols 19-24: Manufacturer Details
        set_cell_value(ws.cell(row=row_idx, column=19), data.get("mfgname", ""))
        set_cell_value(ws.cell(row=row_idx, column=20), data.get("mfgadd1", ""))
        set_cell_value(ws.cell(row=row_idx, column=21), data.get("mfgadd2", ""))
        set_cell_value(ws.cell(row=row_idx, column=22), data.get("mfgadd3", ""))
        set_cell_value(ws.cell(row=row_idx, column=23), data.get("distcd", ""))
        set_cell_value(ws.cell(row=row_idx, column=24), data.get("statecd", ""))

        # Cols 25-29: TESTREPO
        set_cell_value(ws.cell(row=row_idx, column=25), data.get("testrepo_q", "Q"))
        set_cell_value(ws.cell(row=row_idx, column=26), data.get("testrepo_quarter", 1), is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=27), data.get("testrepo_no", ""))
        set_cell_value(ws.cell(row=row_idx, column=28), data.get("testrepo_yyyy", 2026), is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=29), data.get("testrepo_group", ""))

        # Cols 30-32: TESTDATE (DD, MM, YYYY)
        t_dd, t_mm, t_yy = parse_date_parts(data.get("testdate_iso"))
        set_cell_value(ws.cell(row=row_idx, column=30), t_dd, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=31), t_mm, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=32), t_yy, is_integer=True)

        # Col 33: RESULT (Integer 1 or 2), Col 34: NSQ, Col 35: DET1_NSQ
        res_val = data.get("result", 1)
        try:
            res_val = int(res_val) if res_val is not None else 1
        except (ValueError, TypeError):
            res_val = 1
        set_cell_value(ws.cell(row=row_idx, column=33), res_val, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=34), data.get("nsq", ""))
        set_cell_value(ws.cell(row=row_idx, column=35), data.get("det1_nsq", ""))

        # Col 36: FEES (Strict String: e.g. "1390" or "NIL")
        set_cell_value(ws.cell(row=row_idx, column=36), data.get("fees", "NIL"), is_strict_string=True)

        # Col 37: AR REPORT NO
        set_cell_value(ws.cell(row=row_idx, column=37), data.get("ar_report_no", ""))

        # Cols 38-40: AR REPORT DATE (DD, MM, YYYY)
        ar_dd, ar_mm, ar_yy = parse_date_parts(data.get("ar_report_date_iso"))
        set_cell_value(ws.cell(row=row_idx, column=38), ar_dd, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=39), ar_mm, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=40), ar_yy, is_integer=True)

        # Col 41: DI MEMO NO (Strict Text @)
        set_cell_value(ws.cell(row=row_idx, column=41), data.get("di_memo_no", ""), is_strict_string=True)

        # Cols 42-44: DI MEMO DATE (DD, MM, YYYY)
        m_dd, m_mm, m_yy = parse_date_parts(data.get("di_memo_date_iso"))
        set_cell_value(ws.cell(row=row_idx, column=42), m_dd, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=43), m_mm, is_integer=True)
        set_cell_value(ws.cell(row=row_idx, column=44), m_yy, is_integer=True)

        # Col 45: GOVERNMENT ANALYST
        set_cell_value(ws.cell(row=row_idx, column=45), data.get("govt_analyst", ""))

        # Col 46: FEES (Present on PDT sheet only)
        if sheet_name == "PDT":
            set_cell_value(ws.cell(row=row_idx, column=46), data.get("extra_fees", ""), is_strict_string=True)

        cls._finalize_write(wb, path)

    @classmethod
    def delete_excel_row(cls, file_path: Path | str, sheet_name: str, row_idx: int, is_monthly: bool = False):
        """
        Deletes a row from an Excel sheet and re-sequences SR. NO if it's a monthly sheet.
        """
        path = Path(file_path)
        cls._prepare_write(path, f"del_r{row_idx}")

        wb = openpyxl.load_workbook(path)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {path.name}")
        ws = wb[sheet_name]

        ws.delete_rows(row_idx, 1)

        # In monthly sheets, re-sequence the SR. NO column
        if is_monthly:
            sr_col = 1 if sheet_name == "ESIS" else 2
            current_sr = 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=sr_col)
                # If row is non-empty
                if any(ws.cell(row=r, column=c).value is not None for c in range(1, ws.max_column + 1)):
                    cell.value = current_sr
                    cell.data_type = 'n'
                    current_sr += 1

        cls._finalize_write(wb, path)
